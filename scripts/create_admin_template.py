"""Create the blank Wapo Book workbook. Requires openpyxl==3.1.5."""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
ROOT=Path(__file__).resolve().parents[1]
CLASSES=['전사','성기사','사냥꾼','도적','사제','주술사','마법사','흑마법사','드루이드','미확인']
RACES=['인간','드워프','노움','나이트 엘프','드레나이','오크','언데드','타우렌','트롤','블러드 엘프','미확인']
ROLES=['탱커','딜러','힐러','미확인']
HEADERS={
 '모험가 정보':['버튜버명','캐릭터명','레벨','종족','직업','전문화','역할','소속 길드','SOOP 방송국 주소','대표 캐릭터'],
 '길드 정보':['길드명','길드장','전체 길드원 수','길드 소개','최종 확인일'],
 '이전 시즌 참여 기록':['버튜버명','참여 여부','당시 레벨','당시 직업','당시 전문화','확인된 레이드 (쉼표 구분)'],
}
def build(path=ROOT/'downloads/wapobook-admin-template.xlsx'):
 w=Workbook();w.remove(w.active)
 for idx,(name,headers) in enumerate(HEADERS.items()):
  s=w.create_sheet(name);s.append(headers);s.append(['']*len(headers));s.freeze_panes='A2';s.sheet_view.showGridLines=False
  s.row_dimensions[1].height=32
  for c in s[1]:c.font=Font(name='맑은 고딕',bold=True,color='FFF1CF');c.fill=PatternFill('solid',fgColor='493523');c.alignment=Alignment(vertical='center')
  for c in s[2]:c.number_format='@'
  for i,h in enumerate(headers,1):s.column_dimensions[s.cell(1,i).column_letter].width=38 if ('주소' in h or '소개' in h or '레이드' in h) else 22
  t=Table(displayName=['Adventurers','Guilds','PastSeasons'][idx],ref=f'A1:{s.cell(1,len(headers)).column_letter}2');t.tableStyleInfo=TableStyleInfo(name='TableStyleMedium4',showRowStripes=True);s.add_table(t)
 def dropdown(sheet,col,values):
  d=DataValidation(type='list',formula1='"'+','.join(values)+'"',allow_blank=True);d.errorTitle='선택 항목 확인';d.error='목록에서 선택해 주세요.';d.showErrorMessage=True;d.errorStyle='stop';w[sheet].add_data_validation(d);d.add(f'{col}2:{col}1048576')
 for col,values in [('D',RACES),('E',CLASSES),('G',ROLES),('J',['예','아니오'])]:dropdown('모험가 정보',col,values)
 dropdown('이전 시즌 참여 기록','B',['참가','미참가','미확인']);dropdown('이전 시즌 참여 기록','D',CLASSES)
 s=w.create_sheet('입력 방법 안내');s.append(['항목','입력 방법'])
 instructions=[
 ('작성 순서','길드 정보 → 모험가 정보 → 이전 시즌 참여 기록. 빈 양식이며 예시 데이터는 등록되지 않습니다.'),
 ('행 추가','150명 제한이 없습니다. 마지막 셀에서 Tab 또는 행 삽입으로 추가하세요. 드롭다운은 엑셀 마지막 행까지 적용됩니다.'),
 ('필수값','버튜버명 또는 길드명만 필수입니다. 캐릭터·사진·이전 시즌 정보는 모르면 비워 두세요.'),
 ('방송국 주소','계정 ID만 입력하거나 https://www.sooplive.com/station/계정ID 주소를 입력하세요. ch.sooplive.co.kr, ch.sooplive.com, bj.afreecatv.com도 인식합니다.'),
 ('프로필 사진','사진 열은 없습니다. 관리자에서 ID로 프로필 사진 조회를 눌러 원본 이미지 로드를 확인하고 승인합니다. 기존 사진은 덮어쓰지 않습니다.'),
 ('등록 전 확인','관리자 → 엑셀 일괄 등록 → 업로드 → 결과 확인 → 성공 행 병합. 업로드만으로 저장되지 않습니다.'),
 ('중복과 업데이트','기존 SOOP 계정 우선, 다음은 정확히 일치하는 버튜버명으로 연결합니다. 기존 값은 유지하고 빈 항목만 채웁니다.'),
 ('여러 캐릭터','같은 버튜버명과 방송국 주소로 여러 행을 작성하고 캐릭터명은 다르게 입력하세요. 대표 캐릭터는 한 명만 예로 선택하세요.'),
 ('길드 연결','소속 길드는 길드 정보 시트 또는 기존 등록 길드명과 정확히 일치해야 합니다. 없는 길드는 오류로 표시됩니다.'),
 ('레벨','1 이상의 정수만 입력합니다. 최댓값을 추정하지 않습니다. 미확인은 비워 두세요.'),
 ('이전 시즌','참가·미참가·미확인 중 선택. 빈칸은 기존 값을 지우지 않습니다. 레이드 기록은 쉼표로 구분합니다.'),
 ('정보 보존','기존 데이터·사진·즐겨찾기·20인 슬롯을 삭제하거나 초기화하지 않습니다. 수동 수정은 기존 관리자 수정 버튼을 사용하세요.'),
 ('공개 반영 범위','현재 사이트는 브라우저 저장 방식입니다. 병합은 현재 브라우저에 반영됩니다. 전체 방문자 공개는 별도 서버/공개 데이터 배포가 필요합니다.'),
 ('사진 동기화','엑셀 업로드 → ID로 프로필 사진 조회 → 이미지 미리보기 확인 → 병합 승인. 방송 여부는 조회하지 않습니다. 실패하면 수동 사진 등록을 사용하세요.'),
 ('수식','수식 셀·매크로 파일은 허용하지 않습니다. 셀에는 값만 입력하세요. 파일 최대 20MB.'),
 ]
 for row in instructions:s.append(row)
 s.column_dimensions['A'].width=23;s.column_dimensions['B'].width=110;s.freeze_panes='A2'
 for row in s:
  for c in row:c.alignment=Alignment(wrap_text=True,vertical='top');c.font=Font(name='맑은 고딕',size=11)
  s.row_dimensions[row[0].row].height=42
 for c in s[1]:c.fill=PatternFill('solid',fgColor='493523');c.font=Font(bold=True,color='FFF1CF')
 Path(path).parent.mkdir(parents=True,exist_ok=True);w.save(path);print(path)
if __name__=='__main__':build()
